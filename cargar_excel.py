import sqlite3
import pandas as pd
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
from database import init_db

DB_PATH = Path(__file__).resolve().parent / "inventario.db"

def cargar_desde_existencia_bodega(ruta_excel):
    """Carga los productos evaluando el valor calculado de las fórmulas en Excel."""
    init_db()
    
    # Cargar con openpyxl en modo data_only=True para extraer resultados de fórmulas
    wb = load_workbook(ruta_excel, data_only=True)
    
    hoja_nombre = next(
        (h for h in wb.sheetnames if h.strip().upper() == "EXISTENCIA BODEGA"), 
        None
    )
    
    if not hoja_nombre:
        print(f"⚠️ No se encontró la pestaña 'EXISTENCIA BODEGA'. Pestañas disponibles: {wb.sheetnames}")
        wb.close()
        return

    ws = wb[hoja_nombre]
    data = list(ws.values)
    wb.close()

    if not data:
        print("⚠️ La hoja de Excel está vacía.")
        return

    # Convertir a DataFrame de Pandas
    df = pd.DataFrame(data[1:], columns=data[0])
    
    # Limpiar nombres de columnas
    df.columns = [str(col).replace('\n', ' ').replace('\r', '').strip() if col is not None else '' for col in df.columns]
    
    col_codigo = df.columns[0]      # Columna A (Codigo)
    col_descrip = df.columns[1]     # Columna B (Material Description)
    
    # Buscar columna de cantidad: 'CANTIDAD BODEGA' o Columna G (índice 6)
    col_cant_match = [c for c in df.columns if 'CANTIDAD BODEGA' in c.upper()]
    if col_cant_match:
        col_cantidad = col_cant_match[0]
    else:
        col_cantidad = df.columns[6] if len(df.columns) > 6 else None
    
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    cargados = 0
    con_stock = 0
    
    for _, row in df.iterrows():
        cod = str(row[col_codigo]).strip() if pd.notna(row[col_codigo]) else None
        
        if not cod or cod.lower() in ['nan', 'none', '', 'codigo']:
            continue
            
        nombre = str(row[col_descrip]).strip() if pd.notna(row[col_descrip]) else "Sin Descripción"
        
        # Procesar valor numérico calculado de la fórmula en Columna G
        cant = 0
        if col_cantidad and pd.notna(row[col_cantidad]):
            try:
                cant = int(float(row[col_cantidad]))
            except (ValueError, TypeError):
                cant = 0
        
        if cant > 0:
            con_stock += 1

        # Mapeo exacto hacia las columnas de SQLite: codigo, nombre, cantidad
        cursor.execute("""
            INSERT INTO productos (codigo, nombre, cantidad)
            VALUES (?, ?, ?)
            ON CONFLICT(codigo) DO UPDATE SET
                nombre = excluded.nombre,
                cantidad = excluded.cantidad
        """, (cod, nombre, cant))
        
        cargados += 1
        
    conn.commit()
    conn.close()
    print(f"✅ ¡Éxito! Se cargaron {cargados} productos ({con_stock} con stock real) a {DB_PATH.name}")


if __name__ == "__main__":
    archivo_excel = r"C:\Users\1872157.PRODUCTION\OneDrive - LLA\Escritorio\BODEGA - bot_tele_COPIA.xlsx"

    if Path(archivo_excel).exists():
        cargar_desde_existencia_bodega(archivo_excel)
    else:
        print(f"⚠️ No se encontró '{archivo_excel}'. Copia el Excel a la carpeta del bot.")