import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from database import SessionLocal, Producto, init_db


def cargar_desde_existencia_bodega(ruta_excel):
    """Carga los productos evaluando el valor calculado de las fórmulas en Excel."""
    init_db()

    ruta = Path(ruta_excel)
    if not ruta.exists():
        print(f"❌ Error: No se encontró el archivo: '{ruta_excel}'")
        return

    # Cargar con openpyxl en modo data_only=True para extraer resultados de fórmulas
    wb = load_workbook(ruta, data_only=True)

    hoja_nombre = next(
        (h for h in wb.sheetnames if h.strip().upper() == "EXISTENCIA BODEGA"),
        None,
    )

    if not hoja_nombre:
        print(f"⚠️ No se encontró la pestaña 'EXISTENCIA BODEGA'. Pestañas disponibles: {wb.sheetnames}")
        wb.close()
        return

    ws = wb[hoja_nombre]
    data = list(ws.values)
    wb.close()

    if not data or len(data) < 2:
        print("⚠️ La hoja de Excel está vacía o no contiene datos.")
        return

    # Convertir a DataFrame de Pandas
    df = pd.DataFrame(data[1:], columns=data[0])

    # Limpiar nombres de columnas
    df.columns = [
        str(col).replace("\n", " ").replace("\r", "").strip()
        if col is not None
        else ""
        for col in df.columns
    ]

    col_codigo = df.columns[0]  # Columna A (Codigo)
    col_descrip = df.columns[1]  # Columna B (Material Description)

    # Buscar columna de cantidad
    col_cant_match = [c for c in df.columns if "CANTIDAD BODEGA" in c.upper()]
    if col_cant_match:
        col_cantidad = col_cant_match[0]
    else:
        col_cantidad = df.columns[6] if len(df.columns) > 6 else None

    session = SessionLocal()
    cargados = 0
    con_stock = 0

    try:
        for _, row in df.iterrows():
            cod_raw = row[col_codigo]
            if pd.isna(cod_raw):
                continue

            # Formatear el código correctamente sin decimales flotantes (.0)
            if isinstance(cod_raw, float) and cod_raw.is_integer():
                cod = str(int(cod_raw)).strip()
            else:
                cod = str(cod_raw).strip()

            if not cod or cod.lower() in ["nan", "none", "", "codigo"]:
                continue

            nombre = (
                str(row[col_descrip]).strip()
                if pd.notna(row[col_descrip])
                else "Sin Descripción"
            )

            # Procesar valor numérico de la cantidad
            cant = 0
            if col_cantidad and pd.notna(row[col_cantidad]):
                try:
                    cant = int(float(row[col_cantidad]))
                except (ValueError, TypeError):
                    cant = 0

            if cant > 0:
                con_stock += 1

            # Insertar o Actualizar Producto en la BD
            producto = (
                session.query(Producto).filter(Producto.codigo == cod).first()
            )
            if producto:
                producto.nombre = nombre
                producto.cantidad = cant
            else:
                producto = Producto(codigo=cod, nombre=nombre, cantidad=cant)
                session.add(producto)

            cargados += 1

        session.commit()
        print(f"✅ ¡Éxito! Se cargaron/actualizaron {cargados} productos ({con_stock} con stock real).")

    except Exception as e:
        session.rollback()
        print(f"❌ Error durante la carga de materiales: {e}")
    finally:
        session.close()


if __name__ == "__main__":
    print("--- INICIANDO CARGA DE MATERIALES ---")
    excel_local = Path(__file__).resolve().parent / "existencia_bodega.xlsx"
    excel_remoto = Path(
        r"C:\Users\1872157.PRODUCTION\OneDrive - LLA\Escritorio\BODEGA - bot_tele_COPIA.xlsx"
    )

    archivo_excel = excel_local if excel_local.exists() else excel_remoto

    if archivo_excel.exists():
        print(f"Cargando desde: {archivo_excel}")
        cargar_desde_existencia_bodega(archivo_excel)
    else:
        print(f"⚠️ No se encontró el archivo de Excel en ninguna de las rutas especificados.")